
import os
import re
from io import BytesIO

import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl import load_workbook
from groq import Groq
from dotenv import load_dotenv

load_dotenv()

# ============================================================
# AI CONSTRUCTION ASSISTANT
# Source workbook:
#   BHATTI 7 MARLA CPLAN.xlsx
# Source sheet:
#   5.32C-full costr cost-HB
#
# IMPORTANT:
# The calculator reads the named spreadsheet at runtime. Put the
# workbook in the same folder as app.py for automatic loading,
# OR upload it from the sidebar.
# ============================================================

APP_TITLE = "AI Construction Assistant"
SOURCE_FILE = "BHATTI 7 MARLA CPLAN.xlsx"
SOURCE_SHEET = "5.32C-full costr cost-HB"

MARLA_SQFT = 272.25
REFERENCE_FLOOR_SQFT = 1200.0
REFERENCE_LEVELS = 5
REFERENCE_TOTAL_SQFT = REFERENCE_FLOOR_SQFT * REFERENCE_LEVELS


# -----------------------------
# Helpers
# -----------------------------
def money(value):
    try:
        return f"PKR {float(value):,.0f}"
    except Exception:
        return "PKR 0"


def number(value):
    try:
        return float(value)
    except Exception:
        return None


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def infer_category(row_num, description):
    """
    Category mapping follows the structure of the exact
    5.32C-full costr cost-HB worksheet.
    """
    d = description.lower()

    if 11 <= row_num <= 22:
        if row_num <= 16:
            return "Waterproofing & Site Works"
        return "Excavation & Grey Labour"

    if 26 <= row_num <= 31:
        return "Grey Structure"

    if 39 <= row_num <= 66:
        return "Plumbing"

    if 71 <= row_num <= 104:
        return "Electrical"

    if 109 <= row_num <= 137:
        if "door" in d or "window" in d or "wood" in d or "wardrobe" in d or "cabinet" in d:
            return "Doors, Windows & Woodwork"
        if "metal" in d or "gate" in d or "railing" in d or "staircase" in d:
            return "Metal Works"
        return "Finishing"

    if 142 <= row_num <= 162:
        return "Fittings & Fixtures"

    return "Other"


def is_total_or_heading(description):
    d = description.upper()
    if not description:
        return True
    if d.startswith("TOTAL "):
        return True
    headings = [
        "ROOF INSULATION",
        "EXCAVATION & BACKFILL",
        "BRICKS/CEMENT/SAND/KRUSH/IRON",
        "PLUMBING WORKS",
        "ELECTRICAL WORKS",
        "WOOD, METAL",
        "FITTING AND FIXTURE",
        "CONSTRUCTION COST",
        "COLD AND HOT WATER",
        "INSULATION ARMAFLEX",
        "DRAINAGE SYSTEM PIPING",
        "GAS PIPING",
        "PUMPS",
        "AC DRAINAGE SYSTEM",
        "UPVC PIPELINE",
        "SWITCHES AND BACK BOXES",
        "CONDUITS & ACCESSORIES",
        "WIRES & CABLES",
        "LIGHTS, FANS",
        "PAINT AND CEILING",
        "TILE WORK",
        "WOODEN DOORS",
        "WINDOWS",
        "METAL WORKS",
        "WARDROBES & CABINETS",
        "KITCHEN",
        "BATHROOM",
        "GEYSER & MANHOLE",
    ]
    return any(d.startswith(h) for h in headings)


def load_source_workbook(uploaded_file=None):
    """
    Load the exact workbook and exact named worksheet.
    """
    if uploaded_file is not None:
        data = uploaded_file.getvalue()
        wb_values = load_workbook(BytesIO(data), data_only=True)
        wb_formulas = load_workbook(BytesIO(data), data_only=False)
    else:
        if not os.path.exists(SOURCE_FILE):
            return None, None, (
                f"'{SOURCE_FILE}' was not found. Upload the original workbook "
                "from the sidebar or place it beside app.py."
            )
        wb_values = load_workbook(SOURCE_FILE, data_only=True)
        wb_formulas = load_workbook(SOURCE_FILE, data_only=False)

    if SOURCE_SHEET not in wb_values.sheetnames:
        return None, None, (
            f"The workbook does not contain the required sheet "
            f"'{SOURCE_SHEET}'. Available sheets: {', '.join(wb_values.sheetnames)}"
        )

    return wb_values[SOURCE_SHEET], wb_formulas[SOURCE_SHEET], None


def extract_reference_items(ws_values, ws_formulas):
    """
    Extract item descriptions, units, rates, quantities and amounts
    from columns C:G of the exact named worksheet.

    If the workbook has a broken formula (#REF!) in quantity/amount,
    the row is handled safely rather than inventing a value.
    """
    items = []
    broken_rows = []

    for r in range(1, ws_values.max_row + 1):
        description = clean_text(ws_values.cell(r, 3).value)
        unit = clean_text(ws_values.cell(r, 4).value)
        rate = number(ws_values.cell(r, 5).value)
        quantity = number(ws_values.cell(r, 6).value)
        amount = number(ws_values.cell(r, 7).value)

        formula_qty = clean_text(ws_formulas.cell(r, 6).value)
        formula_amount = clean_text(ws_formulas.cell(r, 7).value)

        if not description or is_total_or_heading(description):
            continue

        # Only rows with a usable construction unit are treated as BOQ items.
        if not unit:
            continue

        # Ignore explicit zero rows unless they are useful lump-sum source items.
        if quantity is not None and quantity <= 0 and (amount is None or amount <= 0):
            continue

        # Detect broken formulas in the original spreadsheet.
        if "#REF!" in formula_qty or "#REF!" in formula_amount:
            broken_rows.append({
                "row": r,
                "description": description,
                "formula_quantity": formula_qty,
                "formula_amount": formula_amount,
            })

        # If rate is missing but source amount exists, preserve it as a
        # reference lump-sum cost. This is especially useful for Iron Bars/Sarya.
        lump_sum = False
        if rate is None and amount is not None and amount > 0:
            lump_sum = True
            reference_quantity = 1.0
            reference_rate = amount
        else:
            reference_quantity = quantity if quantity is not None else 0.0
            reference_rate = rate if rate is not None else 0.0

        if reference_quantity <= 0 and not lump_sum:
            continue

        category = infer_category(r, description)

        items.append({
            "source_row": r,
            "category": category,
            "description": description,
            "unit": unit,
            "rate": reference_rate,
            "reference_quantity": reference_quantity,
            "reference_amount": amount if amount is not None else reference_quantity * reference_rate,
            "lump_sum": lump_sum,
        })

    return items, broken_rows


def room_adjusted_quantity(item, base_quantity, bedrooms, washrooms, drawing_rooms,
                           lobby, laundry, basement, floors, total_area):
    """
    Apply conservative residential adjustments to the source quantities.

    The spreadsheet is a commercial-plaza reference, so room-count inputs
    are used only where they naturally affect item counts. Area-scaled
    materials remain area-scaled.
    """
    d = item["description"].lower()
    cat = item["category"]

    # Basement-specific source items.
    if "basement" in d and not basement:
        return 0.0

    # Kitchen fittings: assume one kitchen for a normal residential house.
    if cat == "Fittings & Fixtures":
        if "kitchen sink" in d or "sink mixer" in d or "gully trap" in d:
            return 1.0

        # Bathroom fixtures scale with number of washrooms.
        bathroom_terms = [
            "water closet",
            "vanity basin",
            "basin mixer",
            "muslim shower",
            "shower set",
            "toilet paper holder",
            "beveled edge mirror",
            "soap dish",
        ]
        if any(term in d for term in bathroom_terms):
            return float(max(washrooms, 0))

        if "geyser" in d:
            return float(max(washrooms, 1))

        if "manhole" in d:
            # Keep the source's two-manholes-per-reference layout as an
            # area-scaled quantity, rather than multiplying by washrooms.
            return base_quantity

    # Doors.
    if cat == "Doors, Windows & Woodwork":
        if "room doors" in d:
            return float(max(bedrooms, 0))
        if "washroom" in d and "door" in d:
            return float(max(washrooms, 0))
        if "main entrance" in d and "gf" in d:
            return 1.0
        if "main entrance" in d and basement:
            return base_quantity
        if "wardrobe" in d and "wardrobes" in d:
            return float(max(bedrooms, 0))
        if "kitchen cabinets" in d:
            return 1.0

    # Windows: simple residential assumption.
    if "window" in d:
        if "bathroom" in d:
            return float(max(washrooms, 0))
        if "room windows" in d:
            return float(max(bedrooms * 2, 0))

    # Lights/fans.
    if cat == "Electrical":
        if "ceiling fans" in d:
            return float(max(bedrooms + drawing_rooms + (1 if lobby else 0), 1))
        if "bathroom lights" in d:
            return float(max(washrooms, 0))
        if "kitchen lights" in d:
            return 1.0
        if "drawing room lights" in d:
            return float(max(drawing_rooms, 0))

    # Laundry is only included when selected.
    if "laundry" in d and not laundry:
        return 0.0

    return base_quantity


def calculate_boq(items, sqft_per_floor, floors, bedrooms, washrooms,
                  drawing_rooms, lobby, laundry, basement):
    total_area = sqft_per_floor * floors

    # The named source worksheet is a 30x40 = 1,200 sq.ft reference
    # and contains Basement + GF + 1st + 2nd + Top = five levels.
    scale = total_area / REFERENCE_TOTAL_SQFT

    result = []

    for item in items:
        base_qty = item["reference_quantity"] * scale

        adjusted_qty = room_adjusted_quantity(
            item,
            base_qty,
            bedrooms,
            washrooms,
            drawing_rooms,
            lobby,
            laundry,
            basement,
            floors,
            total_area,
        )

        if adjusted_qty <= 0:
            continue

        amount = adjusted_qty * item["rate"]

        # For lump-sum source rows, rate represents the reference total cost
        # and quantity is a scale multiplier.
        if item["lump_sum"]:
            display_unit = "Reference lump sum"
            display_qty = adjusted_qty
        else:
            display_unit = item["unit"]
            display_qty = adjusted_qty

        result.append({
            "Category": item["category"],
            "Description": item["description"],
            "Unit": display_unit,
            "Quantity": display_qty,
            "Rate (PKR)": item["rate"],
            "Amount (PKR)": amount,
            "Source Row": item["source_row"],
        })

    return pd.DataFrame(result), scale, total_area


def add_subtotals(df):
    if df.empty:
        return pd.DataFrame(columns=["Category", "Subtotal (PKR)"])

    return (
        df.groupby("Category", as_index=False)["Amount (PKR)"]
        .sum()
        .rename(columns={"Amount (PKR)": "Subtotal (PKR)"})
        .sort_values("Subtotal (PKR)", ascending=False)
    )


def make_excel(df, summary_df, inputs):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        df.to_excel(writer, index=False, sheet_name="BOQ")
        summary_df.to_excel(writer, index=False, sheet_name="Cost Summary")

        input_df = pd.DataFrame(
            [{"Input": k, "Value": v} for k, v in inputs.items()]
        )
        input_df.to_excel(writer, index=False, sheet_name="Project Inputs")

        workbook = writer.book
        money_fmt = workbook.add_format({"num_format": '#,##0'})
        qty_fmt = workbook.add_format({"num_format": '#,##0.00'})

        ws = writer.sheets["BOQ"]
        ws.set_column("A:A", 28)
        ws.set_column("B:B", 70)
        ws.set_column("C:C", 20)
        ws.set_column("D:D", 14, qty_fmt)
        ws.set_column("E:F", 16, money_fmt)
        ws.set_column("G:G", 12)

    output.seek(0)
    return output.getvalue()


def groq_client(api_key):
    if not api_key:
        return None
    return Groq(api_key=api_key)


def ask_groq(client, prompt, model):
    response = client.chat.completions.create(
        model=model,
        temperature=0.2,
        max_tokens=1800,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a Pakistan-specific construction cost assistant. "
                    "The Python calculator is authoritative for quantities and costs. "
                    "Do not invent spreadsheet rates. Explain assumptions clearly. "
                    "The source is the worksheet '5.32C-full costr cost-HB' from the "
                    "user's workbook. Costs are estimates and should be verified with "
                    "contractors/suppliers before construction."
                ),
            },
            {"role": "user", "content": prompt},
        ],
    )
    return response.choices[0].message.content


# ============================================================
# Streamlit UI
# ============================================================
st.set_page_config(
    page_title=APP_TITLE,
    page_icon="🏗️",
    layout="wide",
)

st.title("🏗️ AI Construction Assistant")
st.caption(
    "Pakistan-specific residential construction planner using the exact "
    f"'{SOURCE_SHEET}' worksheet from the supplied workbook."
)

# Sidebar
st.sidebar.header("Project Inputs")

uploaded = st.sidebar.file_uploader(
    "Upload source Excel (optional)",
    type=["xlsx", "xlsm"],
    help=(
        f"Use the supplied '{SOURCE_FILE}'. The app specifically reads "
        f"the '{SOURCE_SHEET}' worksheet."
    ),
)

if uploaded is not None:
    ws_values, ws_formulas, load_error = load_source_workbook(uploaded)
else:
    ws_values, ws_formulas, load_error = load_source_workbook()

if load_error:
    st.error(load_error)
    st.info(
        f"Place '{SOURCE_FILE}' in the same folder as app.py or upload it above."
    )
    st.stop()

reference_items, broken_rows = extract_reference_items(ws_values, ws_formulas)

marla = st.sidebar.number_input(
    "Plot size (Marla)",
    min_value=1.0,
    max_value=100.0,
    value=5.0,
    step=0.5,
)

auto_sqft = marla * MARLA_SQFT

use_override = st.sidebar.checkbox(
    "Manually override floor area",
    value=False,
)

if use_override:
    sqft_per_floor = st.sidebar.number_input(
        "Covered area per floor (sq.ft)",
        min_value=100.0,
        max_value=10000.0,
        value=min(auto_sqft, 1200.0),
        step=25.0,
    )
else:
    sqft_per_floor = st.sidebar.number_input(
        "Covered area per floor (sq.ft)",
        min_value=100.0,
        max_value=10000.0,
        value=float(round(auto_sqft, 2)),
        step=25.0,
        disabled=True,
    )

floors = st.sidebar.selectbox(
    "No. of levels / floors",
    options=[1, 2, 3, 4, 5],
    format_func=lambda x: ["Ground", "G+1", "G+2", "G+3", "G+4"][x - 1],
    index=1,
)

bedrooms = st.sidebar.number_input(
    "Bedrooms",
    min_value=0,
    max_value=20,
    value=3,
    step=1,
)

washrooms = st.sidebar.number_input(
    "Washrooms",
    min_value=0,
    max_value=20,
    value=3,
    step=1,
)

drawing_rooms = st.sidebar.number_input(
    "Drawing rooms",
    min_value=0,
    max_value=10,
    value=1,
    step=1,
)

lobby = st.sidebar.checkbox("Lobby", value=True)
laundry = st.sidebar.checkbox("Laundry room", value=False)
basement = st.sidebar.checkbox("Basement", value=False)

st.sidebar.divider()
st.sidebar.header("Groq AI")

api_key = st.sidebar.text_input(
    "Groq API key",
    value=os.getenv("GROQ_API_KEY", ""),
    type="password",
)

model = st.sidebar.text_input(
    "Groq model",
    value="llama-3.3-70b-versatile",
)

# Calculate
boq_df, scale, total_area = calculate_boq(
    reference_items,
    sqft_per_floor,
    floors,
    bedrooms,
    washrooms,
    drawing_rooms,
    lobby,
    laundry,
    basement,
)

summary_df = add_subtotals(boq_df)
grand_total = float(boq_df["Amount (PKR)"].sum()) if not boq_df.empty else 0.0

# Header metrics
c1, c2, c3, c4 = st.columns(4)
c1.metric("Plot size", f"{marla:g} Marla")
c2.metric("Floor area", f"{sqft_per_floor:,.0f} sq.ft")
c3.metric("Total covered area", f"{total_area:,.0f} sq.ft")
c4.metric("Estimated construction cost", money(grand_total))

st.info(
    f"Reference: 30 × 40 = 1,200 sq.ft per level. "
    f"The source worksheet contains 5 reference levels, so the full reference "
    f"area is {REFERENCE_TOTAL_SQFT:,.0f} sq.ft. "
    f"Your project scale factor is {scale:.4f}."
)

if broken_rows:
    with st.expander("⚠️ Source spreadsheet warnings"):
        st.write(
            "The original worksheet contains broken Excel formulas (#REF!). "
            "The app does not invent values for these rows."
        )
        warn_df = pd.DataFrame(broken_rows)
        st.dataframe(warn_df, use_container_width=True, hide_index=True)

# Tabs
tab_boq, tab_ai, tab_summary = st.tabs(
    ["📋 BOQ", "🤖 AI Assistant", "📊 Cost Summary"]
)

with tab_boq:
    st.subheader("Bill of Quantities")

    if boq_df.empty:
        st.warning("No BOQ items were produced.")
    else:
        display_df = boq_df.copy()
        display_df["Quantity"] = display_df["Quantity"].round(2)
        display_df["Rate (PKR)"] = display_df["Rate (PKR)"].round(0)
        display_df["Amount (PKR)"] = display_df["Amount (PKR)"].round(0)

        st.dataframe(
            display_df,
            use_container_width=True,
            hide_index=True,
            column_config={
                "Amount (PKR)": st.column_config.NumberColumn(
                    format="PKR %d"
                ),
                "Rate (PKR)": st.column_config.NumberColumn(
                    format="PKR %d"
                ),
                "Quantity": st.column_config.NumberColumn(
                    format="%.2f"
                ),
            },
        )

        st.subheader("Category Subtotals")
        summary_display = summary_df.copy()
        summary_display["Subtotal (PKR)"] = summary_display["Subtotal (PKR)"].round(0)
        st.dataframe(
            summary_display,
            use_container_width=True,
            hide_index=True,
            column_config={
                "Subtotal (PKR)": st.column_config.NumberColumn(
                    format="PKR %d"
                )
            },
        )

        st.success(f"Grand Total: {money(grand_total)}")

        inputs = {
            "Plot Size (Marla)": marla,
            "Marla conversion (sq.ft)": MARLA_SQFT,
            "Covered Area/Floor (sq.ft)": sqft_per_floor,
            "Number of Floors": floors,
            "Total Covered Area (sq.ft)": total_area,
            "Bedrooms": bedrooms,
            "Washrooms": washrooms,
            "Drawing Rooms": drawing_rooms,
            "Lobby": lobby,
            "Laundry": laundry,
            "Basement": basement,
            "Reference Sheet": SOURCE_SHEET,
            "Scale Factor": scale,
            "Grand Total (PKR)": grand_total,
        }

        excel_bytes = make_excel(boq_df, summary_df, inputs)
        csv_bytes = boq_df.to_csv(index=False).encode("utf-8")

        d1, d2 = st.columns(2)
        with d1:
            st.download_button(
                "⬇️ Download BOQ Excel",
                data=excel_bytes,
                file_name="ai_construction_boq.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        with d2:
            st.download_button(
                "⬇️ Download BOQ CSV",
                data=csv_bytes,
                file_name="ai_construction_boq.csv",
                mime="text/csv",
            )

with tab_ai:
    st.subheader("AI Construction Assistant")

    if not api_key:
        st.warning(
            "Enter your Groq API key in the sidebar to use AI analysis and chat."
        )
    else:
        client = groq_client(api_key)

        if not boq_df.empty:
            boq_context = boq_df[
                ["Category", "Description", "Unit", "Quantity", "Rate (PKR)", "Amount (PKR)"]
            ].to_string(index=False)

            default_prompt = f"""
Analyze this residential construction estimate.

Project:
- Plot: {marla} marla
- Covered area/floor: {sqft_per_floor:,.0f} sq.ft
- Floors: {floors}
- Total covered area: {total_area:,.0f} sq.ft
- Bedrooms: {bedrooms}
- Washrooms: {washrooms}
- Drawing rooms: {drawing_rooms}
- Lobby: {lobby}
- Laundry: {laundry}
- Basement: {basement}

Grand estimate: {money(grand_total)}

BOQ:
{boq_context}

Give:
1. A short cost summary.
2. The three largest cost categories.
3. Important assumptions.
4. Items that should be verified with a contractor.
5. Practical ways to control cost without compromising structural safety.
"""

            if st.button("✨ Generate AI Cost Analysis"):
                with st.spinner("Analyzing the BOQ..."):
                    try:
                        answer = ask_groq(client, default_prompt, model)
                        st.markdown(answer)
                    except Exception as e:
                        st.error(f"Groq error: {e}")

        st.divider()
        st.subheader("Ask a construction question")

        if "chat_messages" not in st.session_state:
            st.session_state.chat_messages = []

        for message in st.session_state.chat_messages:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

        user_question = st.chat_input(
            "Example: Which category is driving my construction cost?"
        )

        if user_question:
            st.session_state.chat_messages.append(
                {"role": "user", "content": user_question}
            )

            context = ""
            if not boq_df.empty:
                context = boq_df[
                    ["Category", "Description", "Unit", "Quantity", "Rate (PKR)", "Amount (PKR)"]
                ].to_string(index=False)

            prompt = f"""
Answer the user's construction question using the calculated BOQ below.

Project total: {money(grand_total)}
Total covered area: {total_area:,.0f} sq.ft
Source worksheet: {SOURCE_SHEET}

BOQ:
{context}

User question:
{user_question}

Rules:
- Do not change or invent calculated quantities.
- If the question asks for a price not present in the BOQ, say that the
  source workbook does not provide that exact price.
- Give practical Pakistan-specific guidance where appropriate.
"""

            with st.chat_message("assistant"):
                with st.spinner("Thinking..."):
                    try:
                        answer = ask_groq(client, prompt, model)
                        st.markdown(answer)
                        st.session_state.chat_messages.append(
                            {"role": "assistant", "content": answer}
                        )
                    except Exception as e:
                        st.error(f"Groq error: {e}")

with tab_summary:
    st.subheader("Cost Breakdown")

    if summary_df.empty:
        st.warning("No cost data available.")
    else:
        fig_bar = px.bar(
            summary_df,
            x="Category",
            y="Subtotal (PKR)",
            title="Cost by Category",
            text_auto=".2s",
        )
        fig_bar.update_layout(
            xaxis_title="Category",
            yaxis_title="PKR",
        )
        st.plotly_chart(fig_bar, use_container_width=True)

        fig_pie = px.pie(
            summary_df,
            names="Category",
            values="Subtotal (PKR)",
            title="Construction Cost Distribution",
        )
        st.plotly_chart(fig_pie, use_container_width=True)

st.caption(
    f"Source: {SOURCE_FILE} → {SOURCE_SHEET}. "
    "This is an estimating tool, not a structural design or contractor quotation."
)
